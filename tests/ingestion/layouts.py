"""Synthetic bank-statement layouts with known ground truth.

Every layout renders the SAME underlying transactions in a different messy
format (Appendix A of the build plan). Tests parse each rendering and compare
the result with the truth, so a regression in any format is caught at once.

These layouts are INVENTED. They imitate common export quirks; they do not
claim to match any real bank. Names are made up. No real data is used.

Use directly:  python -m tests.ingestion.layouts   (prints a summary)
"""

from __future__ import annotations

import csv
import io
import random
from dataclasses import dataclass
from datetime import datetime, timedelta
from decimal import Decimal
from typing import Callable, List, Optional, Tuple

NAMES = ["Rahul Sharma", "Ayesha Khan", "Priya Nair", "Mohammed Irfan", "Sneha Reddy",
         "Karthik Rao", "Fatima Begum", "Arjun Mehta", "Divya Iyer", "Jos\u00e9 Fern\u00e1ndez"]
HANDLES = ["okhdfcbank", "oksbi", "ybl", "okaxis", "paytm"]
VENDORS = ["Stage Rentals", "Print Shop", "Catering Co", "Sound Systems", "Stationery Hub"]
OPENING = Decimal("50000.00")


@dataclass
class Txn:
    dt: datetime
    debit: Optional[Decimal]
    credit: Optional[Decimal]
    balance: Decimal
    ref: Optional[str]  # 12-digit reference present in the narration (None = none)
    name: str
    kind: str  # 'upi' | 'neft'


@dataclass
class Fixture:
    name: str
    filename: str
    data: bytes
    truth: List[Txn]
    has_time: bool = True
    note: str = ""


def build_truth(seed: int = 7, n: int = 80) -> List[Txn]:
    rng = random.Random(seed)
    when = datetime(2026, 10, 1, 9, 0)
    balance = OPENING
    out: List[Txn] = []
    big = [Decimal("150000.00"), Decimal("250000.75")]
    for i in range(n):
        when += timedelta(minutes=rng.randrange(120, 720))
        when = when.replace(second=0)
        if i < 24:  # fixed-fee block: many identical amounts
            credit, debit = Decimal("500.00"), None
        elif rng.random() < 0.6:
            credit = rng.choice([Decimal("250.00"), Decimal("500.00"), Decimal("1200.50"), Decimal("2500.00")] + big)
            debit = None
        else:
            debit = Decimal(rng.randrange(10000, 400000)) / 100
            credit = None
            if debit > balance - 1000:
                debit, credit = None, Decimal("1000.00")
        balance = balance + (credit or Decimal(0)) - (debit or Decimal(0))
        kind = "neft" if i % 9 == 8 else "upi"
        ref = None if kind == "neft" else str(rng.randrange(10**11, 10**12))
        name = rng.choice(NAMES) if credit else rng.choice(VENDORS)
        out.append(Txn(when, debit, credit, balance, ref, name, kind))
    return out


# ---- formatting helpers ------------------------------------------------------


def f2(x: Optional[Decimal]) -> str:
    return "" if x is None else f"{x:.2f}"


def indian(x: Decimal) -> str:
    s = f"{abs(x):.2f}"
    whole, frac = s.split(".")
    if len(whole) > 3:
        head, tail = whole[:-3], whole[-3:]
        parts = []
        while len(head) > 2:
            parts.insert(0, head[-2:])
            head = head[:-2]
        if head:
            parts.insert(0, head)
        whole = ",".join(parts + [tail])
    return ("-" if x < 0 else "") + whole + "." + frac


def narration(t: Txn, style: int = 0) -> str:
    handle = HANDLES[int(t.ref[-1]) % len(HANDLES)] if t.ref else "hdfc"
    if t.kind == "neft":
        return f"NEFT-SBIN{sum(map(ord, t.name)) * 7919 % 10**9:09d}A-{t.name.upper()}-Fest fee"
    if style == 0:
        return f"UPI/{t.ref}/{t.name}/{handle}/Fest fee"
    return f"UPI-{t.name.upper()}-{t.name.split()[0].lower()}@{handle}-HDFC-{t.ref}-Payment"


def _csv(rows: List[List[str]], delimiter: str = ",") -> str:
    buf = io.StringIO()
    csv.writer(buf, delimiter=delimiter, lineterminator="\n").writerows(rows)
    return buf.getvalue()


# ---- the 11 formats ------------------------------------------------------------


def fmt1(truth):  # standard columns, DD/MM/YY
    rows = [["Date", "Narration", "Chq./Ref.No.", "Value Dt", "Withdrawal Amt.", "Deposit Amt.", "Closing Balance"]]
    for t in truth:
        d = t.dt.strftime("%d/%m/%y")
        rows.append([d, narration(t), t.ref or "", d, f2(t.debit), f2(t.credit), f2(t.balance)])
    return Fixture("01_standard", "stmt1.csv", _csv(rows).encode(), truth, has_time=False)


def fmt2(truth):  # one amount column with Dr/Cr suffix
    rows = [["Date", "Narration", "Amount", "Balance"]]
    for t in truth:
        amt = f"{indian(t.credit)} Cr" if t.credit else f"{indian(t.debit)} Dr"
        rows.append([t.dt.strftime("%d/%m/%Y"), narration(t, 1), amt, f"{indian(t.balance)} Cr"])
    return Fixture("02_amount_drcr_suffix", "stmt2.csv", _csv(rows).encode(), truth, has_time=False)


def fmt3(truth):  # one signed amount column
    rows = [["Transaction Date", "Description", "Amount", "Closing Balance"]]
    for t in truth:
        amt = f2(t.credit) if t.credit else "-" + f2(t.debit)
        rows.append([t.dt.strftime("%d-%m-%Y"), narration(t), amt, f2(t.balance)])
    return Fixture("03_signed_amount", "stmt3.csv", _csv(rows).encode(), truth, has_time=False)


def fmt4(truth):  # DD-Mon-YYYY with a separate time column
    rows = [["Date", "Time", "Description", "Debit", "Credit", "Balance"]]
    for t in truth:
        rows.append([t.dt.strftime("%d-%b-%Y"), t.dt.strftime("%H:%M:%S"), narration(t), f2(t.debit), f2(t.credit), f2(t.balance)])
    return Fixture("04_mon_dates_separate_time", "stmt4.csv", _csv(rows).encode(), truth)


def fmt5(truth):  # Indian grouping + currency symbols inside cells
    rows = [["Date", "Narration", "Debit (INR)", "Credit (INR)", "Balance (INR)"]]
    for i, t in enumerate(truth):
        sym = "\u20b9" if i % 2 else "Rs. "
        rows.append([t.dt.strftime("%d %b %Y"), narration(t),
                     sym + indian(t.debit) if t.debit else "", sym + indian(t.credit) if t.credit else "",
                     "\u20b9" + indian(t.balance)])
    return Fixture("05_indian_grouping_symbols", "stmt5.csv", ("\ufeff" + _csv(rows)).encode("utf-8"), truth, has_time=False)


def fmt6(truth):  # junk header block, opening balance row, footer totals
    junk = [["Sample Bank Ltd"], ["Account Statement"], ["Account No: 00112233445566"],
            ["Customer: Fest Society"], ["Address: 12 Example Road, Hyderabad"],
            ["Statement period: 01/10/2026 to 31/10/2026"], []]
    rows = junk + [["Date", "Narration", "Withdrawals", "Deposits", "Balance"],
                   ["01/10/2026", "Opening Balance", "", "", f2(OPENING)]]
    for t in truth:
        rows.append([t.dt.strftime("%d/%m/%Y"), narration(t), f2(t.debit), f2(t.credit), f2(t.balance)])
    td = sum((t.debit or 0 for t in truth), Decimal(0))
    tc = sum((t.credit or 0 for t in truth), Decimal(0))
    rows += [[], ["", "Total", f2(td), f2(tc), ""], ["", "Closing Balance", "", "", f2(truth[-1].balance)],
             ["*** End of statement ***"], ["This is a computer generated statement and needs no signature."]]
    return Fixture("06_junk_header_footer", "stmt6.csv", _csv(rows).encode(), truth, has_time=False)


def fmt7(truth):  # narrations wrapped across two rows (some mid-reference)
    rows = [["Date", "Narration", "Debit", "Credit", "Balance"]]
    for i, t in enumerate(truth):
        n = narration(t)
        cut = 10 if i % 5 == 0 and t.ref else 28  # every 5th breaks INSIDE the 12-digit reference
        first, rest = n[:cut], n[cut:]
        rows.append([t.dt.strftime("%d/%m/%Y"), first, f2(t.debit), f2(t.credit), f2(t.balance)])
        if rest:
            rows.append(["", rest, "", "", ""])
    return Fixture("07_wrapped_narrations", "stmt7.csv", _csv(rows).encode(), truth, has_time=False,
                   note="refs split mid-number on every 5th row")


def fmt8(truth):  # semicolon delimiter, cp1252, decimal comma + dot grouping
    def de(x):
        if x is None:
            return ""
        s = f"{abs(x):,.2f}"
        return s.replace(",", "X").replace(".", ",").replace("X", ".")

    rows = [["Date", "Narration", "Debit", "Credit", "Balance"]]
    for t in truth:
        rows.append([t.dt.strftime("%d/%m/%Y"), narration(t), de(t.debit), de(t.credit), de(t.balance)])
    return Fixture("08_semicolon_cp1252_decimal_comma", "stmt8.csv", _csv(rows, ";").encode("cp1252"), truth, has_time=False)


def fmt9(truth):  # XLSX: merged cells, spacer column, real dates and floats
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Sample Bank Statement"
    ws.merge_cells("A1:F1")
    for col, h in zip("ABDEF", ["Date", "Narration", "Debit", "Credit", "Balance"]):
        ws[f"{col}3"] = h
    ws.merge_cells("B3:C3")  # merged header over a blank spacer column
    r = 4
    for t in truth:
        ws[f"A{r}"] = t.dt
        ws[f"B{r}"] = narration(t)
        if t.debit is not None:
            ws[f"D{r}"] = float(t.debit)
        if t.credit is not None:
            ws[f"E{r}"] = float(t.credit)
        ws[f"F{r}"] = float(t.balance)
        r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return Fixture("09_xlsx_merged_spacer", "stmt9.xlsx", buf.getvalue(), truth)


def fmt10(truth):  # references embedded in several narration patterns
    rows = [["Date", "Narration", "Debit", "Credit", "Balance"]]
    for i, t in enumerate(truth):
        if t.kind == "neft":
            n = f"NEFT-SBIN{i:09d}A-{t.name.upper()}-Fest fee"
        else:
            n = [f"UPI/{t.ref}/{t.name}/okaxis",
                 f"IMPS/P2A/{t.ref}/{t.name}",
                 f"Ref No {t.ref} {t.name}",
                 f"UPI-{t.name.upper()}-x@okaxis-HDFC-{t.ref}-Payment",
                 f"UTR: {t.ref} from {t.name}",
                 f"{t.name} {t.ref}"][i % 6]
        rows.append([t.dt.strftime("%d/%m/%Y"), n, f2(t.debit), f2(t.credit), f2(t.balance)])
    return Fixture("10_embedded_reference_patterns", "stmt10.csv", _csv(rows).encode(), truth, has_time=False)


def fmt11(truth):  # app-style export: newest first, renamed columns, DEBIT/CREDIT type
    rows = [["Date & Time", "Transaction Details", "Type", "Amount (\u20b9)", "Balance (\u20b9)"]]
    for t in reversed(truth):
        rows.append([t.dt.strftime("%d/%m/%Y %H:%M"), narration(t), "CREDIT" if t.credit else "DEBIT",
                     f2(t.credit or t.debit), f2(t.balance)])
    return Fixture("11_app_style_newest_first", "stmt11.csv", ("\ufeff" + _csv(rows)).encode("utf-8"), truth)


# ---- extra stress fixtures --------------------------------------------------------


def fx_inverted_sign(truth):  # signed amount, but positive = money OUT (needs arithmetic repair)
    rows = [["Date", "Details", "Amount", "Balance"]]
    for t in truth:
        amt = f2(t.debit) if t.debit else "-" + f2(t.credit)
        rows.append([t.dt.strftime("%d/%m/%Y"), narration(t), amt, f2(t.balance)])
    return Fixture("X1_inverted_sign_convention", "x1.csv", _csv(rows).encode(), truth, has_time=False)


def fx_no_balance(truth):  # no balance column at all
    rows = [["Date", "Details", "Type", "Amount"]]
    for t in truth:
        rows.append([t.dt.strftime("%d/%m/%Y"), narration(t), "CR" if t.credit else "DR", f2(t.credit or t.debit)])
    return Fixture("X2_no_balance_column", "x2.csv", _csv(rows).encode(), truth, has_time=False)


def fx_headerless(truth):  # no header row at all
    rows = [[t.dt.strftime("%d/%m/%Y"), narration(t), f2(t.debit), f2(t.credit), f2(t.balance)] for t in truth]
    return Fixture("X3_no_header_row", "x3.csv", _csv(rows).encode(), truth, has_time=False)


def fx_tab_text(truth):  # pasted tab-separated text
    rows = [["Date", "Narration", "Debit", "Credit", "Balance"]]
    for t in truth:
        rows.append([t.dt.strftime("%d/%m/%Y"), narration(t), f2(t.debit), f2(t.credit), f2(t.balance)])
    return Fixture("X4_tab_separated", "x4.txt", _csv(rows, "\t").encode(), truth, has_time=False)


TRUTH = build_truth()
CHAOS_BUILDERS: List[Callable] = [fmt1, fmt2, fmt3, fmt4, fmt5, fmt6, fmt7, fmt8, fmt9, fmt10, fmt11]
EXTRA_BUILDERS: List[Callable] = [fx_inverted_sign, fx_no_balance, fx_headerless, fx_tab_text]
CHAOS_FIXTURES: List[Fixture] = [b(TRUTH) for b in CHAOS_BUILDERS]
EXTRA_FIXTURES: List[Fixture] = [b(TRUTH) for b in EXTRA_BUILDERS]
ALL_FIXTURES: List[Fixture] = CHAOS_FIXTURES + EXTRA_FIXTURES


def expected_reference(fx: Fixture, t: Txn, index: int) -> Optional[str]:
    """What the extractor SHOULD return for a transaction in a given fixture."""
    return t.ref


if __name__ == "__main__":  # pragma: no cover
    for fx in ALL_FIXTURES:
        print(f"{fx.name:40s} {len(fx.data):>7} bytes  {fx.filename}")